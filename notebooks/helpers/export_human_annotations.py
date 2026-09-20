"""Export a blinded human-rating workbook for notebook 03; run from repo root."""

from __future__ import annotations

import argparse
import hashlib
from io import BytesIO
import json
from pathlib import Path
import random
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from buli_news.news.annotations.config import (
    DEFAULT_ANNOTATION_CONFIG_PATH,
    INDICATORS,
    load_annotation_config,
)
from buli_news.news.annotations.ollama import build_user_prompt
from buli_news.news.annotations.results import read_existing_annotations
from buli_news.paths import SeasonPaths
from buli_news.storage import read_json


INDICATOR_NAMES = (
    "Sportliche Form", "Personalsituation", "Physische Einsatzbereitschaft",
    "Selbstvertrauen und Motivation",
)
HEADERS = (
    "Taskid", "System Prompt", "User Prompt", "Artikeltitel", "Artikel",
    "Zielteam", *INDICATOR_NAMES,
)
CELL_LIMIT = 32767


def file_sha256(path: Path) -> str:
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def set_text(cell, text: str) -> None:
    """Preserve literal text, including strings that Excel might treat as formulas."""
    if len(text.encode("utf-16-le")) // 2 > CELL_LIMIT:
        raise ValueError(f"Text exceeds Excel cell limit: {cell.coordinate}")
    if ILLEGAL_CHARACTERS_RE.search(text):
        raise ValueError(f"Text contains unsupported XML characters: {cell.coordinate}")
    cell.value = text
    cell.data_type = "s"


def style_sheet(sheet, widths: list[int], row_height: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.freeze_panes = "B2"
    # Sorting individual rows would detach continuation text from its task.
    sheet.auto_filter.ref = None
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet:
        sheet.row_dimensions[row[0].row].height = row_height
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="243746")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F6F8")
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="24495C")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")


def load_sample(paths: SeasonPaths, size: int, seed: int, config) -> tuple:
    """Sample successful productive tasks uniformly without replacement."""
    quality = read_json(paths.news_features_quality)
    provenance = quality["annotation_provenance"]
    if config.sha256 != provenance["annotation_config_sha256"]:
        raise ValueError("Annotation configuration differs from productive report.")
    results = read_existing_annotations(paths.news_annotations)
    eligible = {}
    for result in results:
        if (
            result["annotation_config_sha256"] != config.sha256
            or result["model_digest"] != provenance["model_digest"]
            or result["output_schema_version"]
            != provenance["annotation_output_schema_version"]
        ):
            continue
        task_id = result["task_id"]
        if task_id in eligible:
            raise ValueError("Duplicate productive task_id in results.")
        eligible[task_id] = {
            key: result[key]
            for key in ("request_id", "content_id", "target_team", "match_id", "side")
        }
    del results
    if not 1 <= size <= len(eligible):
        raise ValueError(f"Sample size must be between 1 and {len(eligible)}.")
    sample_ids = random.Random(seed).sample(sorted(eligible), size)
    selected = set(sample_ids)
    found = {}
    with paths.news_annotation_tasks.open(encoding="utf-8-sig") as handle:
        for line in handle:
            if not line.strip():
                continue
            task = json.loads(line)
            task_id = task["task_id"]
            if task_id not in selected:
                continue
            if task_id in found:
                raise ValueError("Duplicate selected task_id in task file.")
            if any(task[key] != value for key, value in eligible[task_id].items()):
                raise ValueError("Task context does not match the productive result.")
            found[task_id] = task
    if set(found) != selected:
        raise ValueError("Not every selected annotation has a corresponding task.")
    return [found[task_id] for task_id in sample_ids], provenance, len(eligible)


def split_cell_text(text: str) -> list[str]:
    """Split losslessly at Excel's UTF-16 cell limit, without breaking characters."""
    chunks = []
    start = 0
    units = 0
    for index, character in enumerate(text):
        width = 2 if ord(character) > 0xFFFF else 1
        if units + width > CELL_LIMIT:
            chunks.append(text[start:index])
            start = index
            units = 0
        units += width
    chunks.append(text[start:])
    return chunks


def build_workbook(
    sample: list[dict[str, Any]], config, notes: list[tuple],
) -> bytes:
    workbook = Workbook()
    tasks = workbook.active
    tasks.title = "Tasks"
    tasks.append(HEADERS)
    # Provenance stays inside the XLSX file, without extra sheets or sidecar files.
    workbook.properties.description = json.dumps(dict(notes), ensure_ascii=False)
    task_blocks = []
    next_row = 2

    for task in sample:
        user_prompt, marker, _ = build_user_prompt(task, config).partition("<article_context>")
        if not marker:
            raise ValueError("Expected article_context marker in production user prompt.")
        values = (
            task["task_id"], config.system_prompt, user_prompt.rstrip(),
            task["title"] or "", task["body"], task["target_team"],
        )
        chunks_by_column = [split_cell_text(value) for value in values]
        height = max(len(chunks) for chunks in chunks_by_column)
        for column, chunks in enumerate(chunks_by_column, start=1):
            for offset, chunk in enumerate(chunks):
                set_text(tasks.cell(next_row + offset, column), chunk)
        task_blocks.append((next_row, height, values))
        next_row += height

    style_sheet(tasks, [25, 42, 48, 42, 85, 25, 26, 26, 28, 34], 110)
    tasks["A1"].comment = Comment(
        "Eine Task beginnt mit ihrer Taskid. Folgende Zeilen ohne Taskid enthalten "
        "gegebenenfalls weiteren Text derselben Task. Nicht einzelne Zeilen sortieren.",
        "Projekt",
    )
    for column, name in enumerate(INDICATOR_NAMES, start=7):
        indicator = INDICATORS[column - 7]
        tasks.cell(1, column).comment = Comment(f"{name} ({indicator})", "Projekt")
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(config.rating_mapping) + '"',
        allow_blank=True,
    )
    validation.showDropDown = False
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    validation.errorTitle = "Ungültiges Rating"
    validation.error = "Bitte eine der sechs Kategorien aus dem Dropdown auswählen."
    tasks.add_data_validation(validation)
    for index, (start_row, height, _) in enumerate(task_blocks):
        for row in tasks.iter_rows(min_row=start_row, max_row=start_row + height - 1):
            for cell in row:
                cell.fill = PatternFill(
                    "solid", fgColor="F3F6F8" if index % 2 == 0 else "FFFFFF",
                )
        for column in range(7, 11):
            tasks.cell(start_row, column).fill = PatternFill("solid", fgColor="FFF2CC")
        validation.add(f"G{start_row}:J{start_row}")
    workbook.active = 0

    buffer = BytesIO()
    workbook.save(buffer)
    payload = buffer.getvalue()
    check = load_workbook(BytesIO(payload))
    for start_row, height, values in task_blocks:
        for column, original in enumerate(values, start=1):
            restored = "".join(
                check["Tasks"].cell(row, column).value or ""
                for row in range(start_row, start_row + height)
            )
            if restored != original:
                raise ValueError(
                    f"Workbook text round-trip failed at row {start_row}, column {column}."
                )
    for row in check["Tasks"].iter_rows(min_row=2, min_col=7, max_col=10):
        if any(cell.value is not None for cell in row):
            raise ValueError("Human rating fields must be blank.")
    check.close()
    return payload


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--season", type=int, default=2025)
    parser.add_argument("--size", type=int, default=200)
    parser.add_argument("--seed", type=int, default=20260906)
    args = parser.parse_args()
    paths = SeasonPaths(args.season)
    destination = (
        Path(__file__).resolve().parent / "human_annotations"
        / f"human_annotations_{args.season}_{args.size}_seed_{args.seed}.xlsx"
    )
    if destination.exists():
        parser.error(f"Workbook already exists; refusing to overwrite manual work: {destination}")
    config = load_annotation_config(DEFAULT_ANNOTATION_CONFIG_PATH)
    sample, provenance, eligible_count = load_sample(paths, args.size, args.seed, config)
    notes = [
        ("Saison", args.season),
        ("Seed", args.seed),
        ("Auswahl", "random.Random(seed).sample(sorted(task_ids), size); ohne Zurücklegen, zufällige Reihenfolge; keine Auswahl nach Ratings."),
        ("Grundgesamtheit", eligible_count),
        ("Stichprobengröße", len(sample)),
        ("Stichprobeneinheit", "Produktiv annotierter Task. Identischer Inhalt in anderem Zielteam-/Request-Kontext bleibt ein eigener Task; keine Inhaltsdeduplizierung."),
        ("Taskid-Reihenfolge SHA-256", hashlib.sha256("\n".join(task["task_id"] for task in sample).encode("utf-8")).hexdigest()),
        ("Konfigurationshash", config.sha256),
        ("Modell-Digest", provenance["model_digest"]),
    ]
    for path in (
        paths.news_annotation_tasks, paths.news_annotations,
        paths.news_features_quality, DEFAULT_ANNOTATION_CONFIG_PATH,
    ):
        notes.append((path.as_posix(), "SHA-256: " + file_sha256(path)))
    payload = build_workbook(sample, config, notes)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("xb") as handle:
        handle.write(payload)
    print(f"Created {len(sample)} tasks with blank human ratings: {destination}")


if __name__ == "__main__":
    main()
