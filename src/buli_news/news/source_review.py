"""Export collected news-source homepages into a review-ready XLSX workbook."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


SOURCE_REVIEW_WORKSHEET = "sources"
SOURCE_REVIEW_COLUMNS = (
    "website_url",
    "article_count",
    "legal_text",
    "review_status",
    "review_date",
)


def export_news_source_review(
    raw_dir: Path,
    output_path: Path,
    *,
    overwrite: bool = False,
) -> list[tuple[str, int]]:
    """Extract source counts and write the workbook used for legal review."""
    homepage_counts = extract_homepage_counts(raw_dir)
    write_source_review_workbook(
        output_path=output_path,
        homepage_counts=homepage_counts,
        overwrite=overwrite,
    )
    return homepage_counts


def extract_homepage_counts(raw_dir: Path) -> list[tuple[str, int]]:
    """Extract normalized source homepage URLs from raw news responses."""
    response_paths = sorted(raw_dir.glob("*.json"))
    if not response_paths:
        msg = f"No raw news response JSON files found in {raw_dir}."
        raise ValueError(msg)

    counts: Counter[str] = Counter()
    for path in response_paths:
        response = read_raw_response(path)
        articles = response.get("articles")
        if not isinstance(articles, dict):
            msg = f"Raw news response {path} has no articles object."
            raise ValueError(msg)
        results = articles.get("results")
        if not isinstance(results, list):
            msg = f"Raw news response {path} has no articles.results list."
            raise ValueError(msg)

        for article_index, article in enumerate(results, start=1):
            if not isinstance(article, dict):
                msg = (
                    f"Raw news response {path} article {article_index} "
                    "must be an object."
                )
                raise ValueError(msg)
            homepage_url = get_homepage_url(article)
            if homepage_url is None:
                msg = (
                    f"Raw news response {path} article {article_index} "
                    "has no valid source URL."
                )
                raise ValueError(msg)
            counts[homepage_url] += 1

    if not counts:
        msg = f"Raw news responses in {raw_dir} contain no articles."
        raise ValueError(msg)
    return sorted(counts.items(), key=lambda item: (-item[1], item[0]))


def read_raw_response(path: Path) -> dict:
    """Read one raw response and require its top-level JSON object."""
    with path.open("r", encoding="utf-8") as file:
        response = json.load(file)
    if not isinstance(response, dict):
        msg = f"Raw news response {path} must contain a JSON object."
        raise ValueError(msg)
    return response


def get_homepage_url(article: dict) -> str | None:
    """Return the normalized homepage URL for one article source."""
    homepage_url = normalize_homepage_url(article.get("url"))
    if homepage_url is not None:
        return homepage_url

    source = article.get("source")
    if isinstance(source, dict):
        return normalize_homepage_url(source.get("uri"))
    return None


def normalize_homepage_url(value: object) -> str | None:
    """Normalize an article URL or source URI to an HTTPS homepage URL."""
    if not isinstance(value, str):
        return None
    stripped = value.strip()
    if not stripped:
        return None

    candidate = stripped if "://" in stripped else f"https://{stripped}"
    try:
        parsed = urlsplit(candidate)
    except ValueError:
        return None
    if parsed.scheme not in {"http", "https"} or parsed.hostname is None:
        return None

    host = parsed.hostname.lower().rstrip(".")
    if host.startswith("www."):
        host = host[4:]
    if not host or any(character.isspace() for character in host):
        return None
    return f"https://{host}"


def write_source_review_workbook(
    output_path: Path,
    homepage_counts: list[tuple[str, int]],
    *,
    overwrite: bool = False,
) -> None:
    """Write the fixed source-review schema as a styled XLSX workbook."""
    if output_path.exists() and not overwrite:
        msg = (
            f"Output workbook already exists: {output_path}. "
            "Use --overwrite only if its manual review data may be replaced."
        )
        raise ValueError(msg)
    if not homepage_counts:
        msg = "Cannot write a source-review workbook without source rows."
        raise ValueError(msg)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SOURCE_REVIEW_WORKSHEET
    worksheet.append(SOURCE_REVIEW_COLUMNS)
    for homepage_url, article_count in homepage_counts:
        worksheet.append((homepage_url, article_count))
        worksheet.cell(worksheet.max_row, 1).alignment = Alignment(vertical="top")
        worksheet.cell(worksheet.max_row, 2).alignment = Alignment(horizontal="right")

    worksheet.freeze_panes = "A2"
    worksheet.row_dimensions[1].height = 22
    for column, width in zip("ABCDE", (34, 14, 70, 32, 14)):
        worksheet.column_dimensions[column].width = width
    for cell in worksheet[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="FF1F4E78")
        cell.alignment = Alignment(vertical="center")

    table = Table(displayName="SourceReview", ref=f"A1:E{worksheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
