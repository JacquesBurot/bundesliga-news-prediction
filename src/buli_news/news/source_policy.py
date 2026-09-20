"""Build a versioned news-source policy from a reviewed XLSX workbook."""

from __future__ import annotations

import hashlib
import re
import zipfile
from collections import Counter
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from openpyxl import load_workbook

from buli_news.news.source_review import (
    SOURCE_REVIEW_COLUMNS,
    SOURCE_REVIEW_WORKSHEET,
)


REVIEW_STATUS_DECISIONS = {
    "complete_no_ml_clause": "include",
    "complete_explicit_ml_clause": "exclude",
}
URL_PATTERN = re.compile(r"https?://[^\s<>]+")
EXCEL_DATE_EPOCH = datetime(1899, 12, 30)


def build_news_source_policy(
    workbook_path: Path,
    season: int,
    policy_version: int = 1,
) -> dict[str, Any]:
    """Convert one reviewed source workbook into a deterministic policy."""
    if isinstance(season, bool) or not isinstance(season, int):
        msg = "Season must be an integer start year."
        raise ValueError(msg)
    if isinstance(policy_version, bool) or not isinstance(policy_version, int):
        msg = "Policy version must be an integer."
        raise ValueError(msg)
    if policy_version < 1:
        msg = "Policy version must be greater than zero."
        raise ValueError(msg)

    rows = read_source_review_workbook(workbook_path)
    rules = [build_policy_rule(row, row_number) for row_number, row in rows]
    rules.sort(key=lambda rule: rule["host"])
    validate_unique_hosts(rules)

    decision_counts = Counter(rule["decision"] for rule in rules)
    review_dates = sorted({rule["review_date"] for rule in rules})
    policy = {
        "schema_version": 1,
        "policy_id": f"news_source_policy_{season}_v{policy_version}",
        "purpose": (
            "Control which collected news sources may enter local-LLM feature "
            "extraction and downstream match-outcome modeling."
        ),
        "source_review": {
            "workbook_path": workbook_path.as_posix(),
            "workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
            "worksheet": SOURCE_REVIEW_WORKSHEET,
            "reviewed_source_count": len(rules),
            "review_date_min": min(review_dates),
            "review_date_max": max(review_dates),
        },
        "decision_policy": {
            "unknown_host_decision": "exclude",
            "host_matching": (
                "Normalize the candidate hostname to lowercase, remove one "
                "leading 'www.' label, and require an exact match with rule.host."
            ),
            "include_status": "complete_no_ml_clause",
            "exclude_status": "complete_explicit_ml_clause",
            "interpretation": (
                "The policy reproduces the review decisions recorded in the "
                "source workbook. It does not independently determine legal "
                "permission."
            ),
        },
        "summary": {
            "include_source_count": decision_counts["include"],
            "exclude_source_count": decision_counts["exclude"],
            "unreviewed_source_count": 0,
            "reviewed_article_occurrence_count": sum(
                rule["article_count"] for rule in rules
            ),
            "included_article_occurrence_count_before_deduplication": sum(
                rule["article_count"]
                for rule in rules
                if rule["decision"] == "include"
            ),
            "excluded_article_occurrence_count_before_deduplication": sum(
                rule["article_count"]
                for rule in rules
                if rule["decision"] == "exclude"
            ),
        },
        "rules": rules,
    }
    validate_news_source_policy(policy)
    return policy


def read_source_review_workbook(
    workbook_path: Path,
) -> list[tuple[int, dict[str, str]]]:
    """Read and validate the fixed source-review worksheet with openpyxl."""
    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
        )
    except zipfile.BadZipFile as exc:
        msg = f"Source review workbook is not a valid XLSX file: {workbook_path}."
        raise ValueError(msg) from exc
    except KeyError as exc:
        msg = f"Source review workbook is missing required XLSX content: {exc}."
        raise ValueError(msg) from exc

    try:
        if SOURCE_REVIEW_WORKSHEET not in workbook.sheetnames:
            msg = f"Source review workbook has no worksheet {SOURCE_REVIEW_WORKSHEET!r}."
            raise ValueError(msg)
        worksheet = workbook[SOURCE_REVIEW_WORKSHEET]
        rows = worksheet.iter_rows(
            min_row=1,
            max_col=max(len(SOURCE_REVIEW_COLUMNS), worksheet.max_column or 0),
            values_only=True,
        )
        header_values = tuple(review_cell_text(value) for value in next(rows, ()))
        headers = header_values[:len(SOURCE_REVIEW_COLUMNS)]
        if headers != SOURCE_REVIEW_COLUMNS:
            msg = (
                "Source review worksheet must have the exact columns "
                f"{list(SOURCE_REVIEW_COLUMNS)} in row 1, got {list(headers)}."
            )
            raise ValueError(msg)
        extra_headers = [
            value
            for value in header_values[len(SOURCE_REVIEW_COLUMNS):]
            if value
        ]
        if extra_headers:
            msg = f"Source review worksheet has unexpected columns: {extra_headers}."
            raise ValueError(msg)

        records = []
        for row_number, row in enumerate(rows, start=2):
            values = tuple(review_cell_text(value) for value in row)
            if not any(values):
                continue
            if any(values[len(SOURCE_REVIEW_COLUMNS):]):
                msg = f"Source review row {row_number} has values outside columns A-E."
                raise ValueError(msg)
            record = dict(zip(SOURCE_REVIEW_COLUMNS, values))
            records.append((row_number, record))

        if not records:
            msg = "Source review worksheet contains no source rows."
            raise ValueError(msg)
        return records
    finally:
        workbook.close()


def review_cell_text(value: object) -> str:
    """Normalize empty cells and Excel dates for the review-field validators."""
    if value is None:
        return ""
    if isinstance(value, datetime) and value.time() == time():
        value = value.date()
    return str(value).strip()


def build_policy_rule(row: dict[str, str], row_number: int) -> dict[str, Any]:
    """Validate one reviewed source row and convert it to one policy rule."""
    website_url = require_text(row, "website_url", row_number)
    host = normalize_reviewed_website_url(website_url, row_number)
    article_count = parse_positive_integer(row, "article_count", row_number)
    review_status = require_text(row, "review_status", row_number)
    try:
        decision = REVIEW_STATUS_DECISIONS[review_status]
    except KeyError as exc:
        msg = (
            f"Source review row {row_number} has unsupported review_status "
            f"{review_status!r}."
        )
        raise ValueError(msg) from exc

    legal_text = row["legal_text"].strip()
    evidence_urls = extract_evidence_urls(legal_text)
    if decision == "exclude" and not legal_text:
        msg = f"Excluded source review row {row_number} needs legal_text."
        raise ValueError(msg)
    if decision == "exclude" and not evidence_urls:
        msg = f"Excluded source review row {row_number} needs an evidence URL."
        raise ValueError(msg)
    if decision == "include" and legal_text:
        msg = (
            f"Included source review row {row_number} must leave legal_text empty."
        )
        raise ValueError(msg)

    return {
        "host": host,
        "website_url": website_url,
        "match_type": "exact_host",
        "decision": decision,
        "review_status": review_status,
        "review_date": parse_review_date(row["review_date"], row_number),
        "article_count": article_count,
        "evidence_urls": evidence_urls,
    }


def normalize_reviewed_website_url(value: str, row_number: int) -> str:
    """Validate one reviewed homepage URL and return its normalized host."""
    parsed = urlsplit(value)
    if parsed.scheme not in {"http", "https"} or parsed.hostname is None:
        msg = f"Source review row {row_number} has invalid website_url {value!r}."
        raise ValueError(msg)
    if parsed.username or parsed.password or parsed.port:
        msg = (
            f"Source review row {row_number} website_url must not contain "
            "user info or a port."
        )
        raise ValueError(msg)
    if parsed.path not in {"", "/"} or parsed.query or parsed.fragment:
        msg = f"Source review row {row_number} website_url must be a homepage URL."
        raise ValueError(msg)

    host = parsed.hostname.lower().rstrip(".")
    if host.startswith("www."):
        host = host[4:]
    if not host:
        msg = f"Source review row {row_number} has an empty website host."
        raise ValueError(msg)
    return host


def parse_positive_integer(row: dict[str, str], key: str, row_number: int) -> int:
    """Parse one required positive whole-number cell."""
    value = require_text(row, key, row_number)
    try:
        numeric_value = float(value)
    except ValueError as exc:
        msg = f"Source review row {row_number} field {key!r} must be numeric."
        raise ValueError(msg) from exc
    if not numeric_value.is_integer() or numeric_value <= 0:
        msg = f"Source review row {row_number} field {key!r} must be positive integer."
        raise ValueError(msg)
    return int(numeric_value)


def parse_review_date(value: str, row_number: int) -> str:
    """Parse an ISO date or an Excel serial date and return YYYY-MM-DD."""
    stripped = value.strip()
    try:
        return date.fromisoformat(stripped).isoformat()
    except ValueError:
        pass

    try:
        numeric_value = float(stripped)
    except ValueError as exc:
        msg = f"Source review row {row_number} has invalid review_date {value!r}."
        raise ValueError(msg) from exc
    if not numeric_value.is_integer() or numeric_value <= 0:
        msg = f"Source review row {row_number} has invalid review_date {value!r}."
        raise ValueError(msg)
    return (EXCEL_DATE_EPOCH + timedelta(days=int(numeric_value))).date().isoformat()


def extract_evidence_urls(legal_text: str) -> list[str]:
    """Extract stable, de-duplicated evidence URLs from reviewed legal text."""
    urls = [url.rstrip(".,);\"'") for url in URL_PATTERN.findall(legal_text)]
    return list(dict.fromkeys(urls))


def validate_unique_hosts(rules: list[dict[str, Any]]) -> None:
    """Reject duplicate normalized hosts in the reviewed workbook."""
    counts = Counter(rule["host"] for rule in rules)
    duplicate_hosts = sorted(host for host, count in counts.items() if count > 1)
    if duplicate_hosts:
        msg = f"Source review workbook contains duplicate hosts: {duplicate_hosts}."
        raise ValueError(msg)


def validate_news_source_policy(policy: dict[str, Any]) -> None:
    """Validate internal counts and ordering in one generated policy."""
    rules = policy["rules"]
    hosts = [rule["host"] for rule in rules]
    if hosts != sorted(hosts):
        msg = "News source policy rules must be sorted by host."
        raise ValueError(msg)
    validate_unique_hosts(rules)

    summary = policy["summary"]
    decision_counts = Counter(rule["decision"] for rule in rules)
    expected_summary = {
        "include_source_count": decision_counts["include"],
        "exclude_source_count": decision_counts["exclude"],
        "unreviewed_source_count": 0,
        "reviewed_article_occurrence_count": sum(
            rule["article_count"] for rule in rules
        ),
        "included_article_occurrence_count_before_deduplication": sum(
            rule["article_count"]
            for rule in rules
            if rule["decision"] == "include"
        ),
        "excluded_article_occurrence_count_before_deduplication": sum(
            rule["article_count"]
            for rule in rules
            if rule["decision"] == "exclude"
        ),
    }
    if summary != expected_summary:
        msg = "News source policy summary does not match its rules."
        raise ValueError(msg)


def require_text(row: dict[str, str], key: str, row_number: int) -> str:
    """Return one required, non-empty reviewed source value."""
    value = row.get(key, "").strip()
    if not value:
        msg = f"Source review row {row_number} needs non-empty field {key!r}."
        raise ValueError(msg)
    return value
